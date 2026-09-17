#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
全社印刷実績の一覧（稼動日報から作った Excel）を母集団にして、通し数 3,000 以下の案件に
帳票 PDF（製造指示書など）から読んだ 仕上りサイズ・加工・用紙 を付け、月別シートの Excel 資料にする。

    python build_final.py --base 2026年9月_全社_印刷実績.xlsx --textdir <pdftext フォルダ> \
                          --out 平版印刷_オンデマンド移行検討_3000通し以下_資料.xlsx [--max-pass 3000]

帳票が無い行は、次のどちらかで埋められる（どちらも「元の帳票」列と地の色で見分けられる）。

    --koutei-csv 仕上りサイズ_koutei取得結果.csv   koutei の DB から取った値（確かな値）
    --guess-reprint                              再版の帳票からの推定（確かな値ではない）

* 母集団の列: 営業部, 管理番号, ｸﾗｲｱﾝﾄ名, 品名, 今年の動向, 無しの場合の代替対策, 対策通し数, 営業担当ｺｰﾄﾞ, 印刷日, 色数, 通し数
  （見出し行が途中に何度も入っているので、管理番号が数字の行だけを明細として読む）
* 月は印刷日の月。同じ管理番号が複数行ある月はそのまま複数行出す（稼動日報の実績どおり）
* 帳票が無い管理番号は、帳票の列が空のまま並ぶ
"""

import argparse
import collections
import datetime as dt
import io
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_from_pdf import load_docs, merge  # noqa: E402
from export_finish_size import judge_a3_all, norm_order  # noqa: E402

FONT_NAME = "Meiryo"
FONT = Font(name=FONT_NAME, size=9)
FONT_BOLD = Font(name=FONT_NAME, size=9, bold=True)
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
FONT_NOTE = Font(name=FONT_NAME, size=8, color="595959")
FONT_HEADER = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
FONT_LINK = Font(name=FONT_NAME, size=9, color="0563C1", underline="single")
FILL_HEADER = PatternFill("solid", fgColor="1F4E78")
FILL_HEADER_K = PatternFill("solid", fgColor="2E75B6")
FILL_BAND = PatternFill("solid", fgColor="F2F2F2")
FILL_A3 = {"○": PatternFill("solid", fgColor="C6EFCE"), "×": PatternFill("solid", fgColor="F8CBAD"),
           "不明": PatternFill("solid", fgColor="FFF2CC")}
FILL_NONE = PatternFill("solid", fgColor="E7E6E6")
FILL_DB = PatternFill("solid", fgColor="DDEBF7")      # koutei の DB から埋めた列
FILL_GUESS = PatternFill("solid", fgColor="E4DFEC")   # 再版元からの推定で埋めた列
SRC_DOC, SRC_DB, SRC_GUESS = "帳票", "koutei(DB)", "再版推定"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALIGN_WRAP = Alignment(vertical="top", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="top")

BASE_COLS = [("営業部", 13), ("管理番号", 10), ("ｸﾗｲｱﾝﾄ名", 20), ("品名", 26), ("営業担当\nｺｰﾄﾞ", 7), ("印刷日", 11), ("色数", 6), ("通し数", 8)]
DOC_COLS = [("仕上りサイズ", 14), ("A3以下", 7), ("加工内容", 26), ("内外作", 8), ("加工所／委託先", 14),
            ("用紙銘柄", 18), ("用紙規格", 10), ("連量", 8), ("総頁数", 6), ("受注数量", 8), ("備考案", 34),
            ("元の帳票", 12), ("帳票\nリンク", 7), ("加工欄の原文", 40)]
COLUMNS = [("No", 5)] + BASE_COLS + DOC_COLS
N_BASE = 1 + len(BASE_COLS)


def toi(v):
    try:
        return int(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def read_base(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        if r is None or len(r) < 11:
            continue
        key, _ = norm_order(r[1])
        if not key or str(r[1]).strip() == "管理番号":
            continue
        rows.append({
            "key": key, "営業部": r[0] or "", "管理番号": str(r[1]).strip(), "ｸﾗｲｱﾝﾄ名": r[2] or "", "品名": r[3] or "",
            "営業担当ｺｰﾄﾞ": r[7] or "", "印刷日": str(r[8] or "")[:10], "色数": r[9] or "", "通し数": toi(r[10]),
        })
    return rows


EMPTY_DOC = {"仕上りサイズ": "", "A3以下": "不明", "加工内容": "", "内外作": "", "加工所": "",
             "用紙銘柄": "", "用紙規格": "", "連量": "", "総頁数": None, "受注数量": None,
             "加工原文": "", "url": ""}


# 帳票の欄が空のとき DB の値で埋める項目（帳票にある値は上書きしない）
MERGE_FIELDS = ["仕上りサイズ", "加工内容", "内外作", "加工所", "用紙銘柄", "用紙規格", "連量"]

KOUTEI_COLS = ["受注番号", "仕上りサイズ", "A3以下", "加工内容", "内外作区分", "委託先名", "用紙銘柄", "用紙規格", "斤量"]


def clean_cell(v):
    """CSV の 1 マス。Excel が受け付けない制御文字を外し、前後の空白を落とす。

    改行・復帰・タブは Excel も openpyxl も受け付けるので残す（加工内容は
    帳票側が改行で区切っており、DB 由来だけ語がつながると読めなくなる）。
    """
    return ILLEGAL_CHARACTERS_RE.sub("", v or "").strip()


def judge_cell(a3, size):
    """CSV の「A3以下」列。○ × 不明 以外が入っていたら仕上りサイズから出し直す。

    仕上りサイズは複数部品をスラッシュでつないだ 1 つの文字列で来るので、
    必ず分けてから渡す（1 部品でも A3 より大きければ × にするため）。
    区切りは対の道具が出す「 / 」に限らず、全角や空白なしでも分ける。
    """
    if a3 in ("○", "×", "不明"):
        return a3
    parts = [x.strip() for x in re.split(r"\s*[/／]\s*", size) if x.strip()]
    return judge_a3_all(parts) if parts else "不明"


def read_koutei_csv(path):
    """export_finish_size.py が出した CSV を読み、帳票と同じ形にして返す。

    列: 受注番号, 仕上りサイズ, A3以下, 加工内容, 内外作区分, 委託先名, 用紙銘柄, 用紙規格, 斤量
    中身が全部空の行は入れない（「DB にも無かった」と「埋まった」を混ぜないため）。
    Excel で開いて保存し直した CP932 の CSV も読む。
    """
    import csv
    text = None
    for enc in ("utf-8-sig", "cp932"):
        try:
            with open(path, encoding=enc, newline="") as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        sys.exit(f"{path} の文字コードが読めません。UTF-8 か CP932 にしてください")
    recs = list(csv.DictReader(io.StringIO(text)))
    head = list(recs[0]) if recs else []
    if recs and "受注番号" not in head:
        sys.exit(f"{path} に「受注番号」の列がありません（export_finish_size.py の出力を渡してください）: {head[:6]}")
    lack = [c for c in KOUTEI_COLS if recs and c not in head]
    if lack:
        print(f"※ {os.path.basename(path)} に無い列があります（空のまま進みます）: {'、'.join(lack)}")
    out, dup = {}, []
    for rec in recs:
        key, _ = norm_order(rec.get("受注番号"))
        if not key:
            continue
        g = lambda k: clean_cell(rec.get(k))  # noqa: E731
        vals = [g(k) for k in ("仕上りサイズ", "加工内容", "内外作区分", "委託先名", "用紙銘柄", "用紙規格", "斤量")]
        if not any(vals):
            continue
        size = g("仕上りサイズ")
        if key in out:          # 先勝ち（make_excel.py と同じ）
            dup.append(key)
            continue
        d = dict(EMPTY_DOC)
        d.update({
            "key": key, "仕上りサイズ": size, "A3以下": judge_cell(g("A3以下"), size),
            "加工内容": g("加工内容"), "内外作": g("内外作区分"), "加工所": g("委託先名"),
            "用紙銘柄": g("用紙銘柄"), "用紙規格": g("用紙規格"), "連量": g("斤量"),
            "帳票": [SRC_DB], "_src": SRC_DB,
        })
        out[key] = d
    if dup:
        print(f"※ {os.path.basename(path)} に同じ受注番号が {len(dup)} 件重なっていました"
              f"（先の行を採用）: {'、'.join(sorted(set(dup))[:5])}")
    return out


def prev_key(d):
    """帳票 1 通が指す「前回受注番号」を正規化して返す。"""
    for one in d.get("docs", [d]):
        k, _ = norm_order(one.get("前回受注番号"))
        if k:
            return k
    return ""


def guess_from_reprint(rows):
    """再版（前回受注番号つき）の帳票から、その「前回受注番号」の仕様を推定する。

    あくまで推定。元も再版も帳票がある 473 組で確かめたところ、両方に仕上りサイズが
    ある 469 組では A3 以下の判定が 96.4% 一致するが、仕上りサイズそのものは 82%、
    加工内容は 67% しか一致しない（連番の別部品を前回受注番号に入れている帳票が
    混ざるため）。使うときは必ず出典を見せること。
    """
    cand = collections.defaultdict(list)
    for r in rows:
        p = prev_key(r)
        if p and p != r["key"]:
            cand[p].append(r)
    out = {}
    for p, rs in cand.items():
        src = sorted(rs, key=lambda r: (r.get("createdTime") or "", r["key"]))[0]
        d = dict(src)
        d.update({"key": p, "帳票": [f"{SRC_GUESS} {src['key']}"], "_src": SRC_GUESS,
                  "url": src.get("url", ""), "総頁数": None, "受注数量": None})
        out[p] = d
    return out


def remark(d):
    """備考欄の文面の案。例: 外注（松岡製本）：ミシン(筋) ／ 内作（第二工場）：中綴じ12P"""
    if not d:
        return ""
    kind, site, cont = d["内外作"], d["加工所"], d["加工内容"]
    if not (kind or site or cont):
        return ""
    head = kind or ""
    if site:
        head = f"{head}（{site}）" if head else f"（{site}）"
    return f"{head}：{cont}" if head and cont else (head or cont)


def style_header(ws, hrow, n_base, n_all):
    for c in range(1, n_all + 1):
        cell = ws.cell(hrow, c)
        cell.font, cell.alignment, cell.border = FONT_HEADER, ALIGN_CENTER, BORDER
        cell.fill = FILL_HEADER if c <= n_base else FILL_HEADER_K
    ws.row_dimensions[hrow].height = 30


def write_month(wb, title, rows, docs, note):
    ws = wb.create_sheet(title)
    ws["A1"] = f"平版印刷 → オンデマンド移行検討　{title}　通し数 3,000 以下"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = note
    ws["A2"].font = FONT_NOTE
    hrow = 4
    for i, (h, w) in enumerate(COLUMNS, 1):
        ws.cell(hrow, i, h)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, hrow, N_BASE, len(COLUMNS))
    r = hrow
    for n, x in enumerate(rows, 1):
        r += 1
        d = docs.get(x["key"])
        vals = [n, x["営業部"], x["管理番号"], x["ｸﾗｲｱﾝﾄ名"], x["品名"], x["営業担当ｺｰﾄﾞ"], x["印刷日"], x["色数"], x["通し数"]]
        if d:
            vals += [d["仕上りサイズ"], d["A3以下"], d["加工内容"], d["内外作"], d["加工所"], d["用紙銘柄"], d["用紙規格"], d["連量"],
                     d["総頁数"], d["受注数量"], remark(d), "・".join(d["帳票"]),
                     "開く" if d.get("url") else "", d["加工原文"]]
        else:
            vals += ["", "帳票なし"] + [""] * (len(DOC_COLS) - 2)
        band = n % 2 == 0
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border, c.alignment = FONT, BORDER, ALIGN_WRAP
            if band:
                c.fill = FILL_BAND
        src = d.get("_src", SRC_DOC) if d else ""
        if src in (SRC_DB, SRC_GUESS):      # 帳票以外から埋めた列は地の色で分ける
            fill = FILL_DB if src == SRC_DB else FILL_GUESS
            for i in range(N_BASE + 1, len(COLUMNS) + 1):
                if i != N_BASE + 2:          # A3以下だけは判定の色を残す
                    ws.cell(r, i).fill = fill
        a3 = ws.cell(r, N_BASE + 2)
        if a3.value in FILL_A3:
            a3.fill, a3.alignment = FILL_A3[a3.value], ALIGN_CENTER
        elif a3.value == "帳票なし":
            a3.fill, a3.alignment, a3.font = FILL_NONE, ALIGN_CENTER, FONT_NOTE
        if d and d.get("url"):
            link = ws.cell(r, N_BASE + len(DOC_COLS) - 1)
            link.hyperlink, link.font = d["url"], FONT_LINK
        for i in (9, N_BASE + 9, N_BASE + 10):
            ws.cell(r, i).number_format = "#,##0"
            ws.cell(r, i).alignment = ALIGN_RIGHT
    ws.freeze_panes = ws.cell(hrow + 1, 4)
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(COLUMNS))}{max(r, hrow + 1)}"
    ws.print_title_rows = f"{hrow}:{hrow}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "&P / &N"
    return hrow + 1, r


CONFLICT_COLS = [("No", 5), ("月", 9), ("営業部", 13), ("管理番号", 10), ("ｸﾗｲｱﾝﾄ名", 20), ("品名", 26), ("通し数", 8),
                 ("帳票の仕上りサイズ", 20), ("帳票\n判定", 7), ("DB の仕上りサイズ", 26), ("DB\n判定", 7),
                 ("採用", 7), ("帳票\nリンク", 7)]


def write_conflicts(wb, rows, taken):
    """帳票と DB で A3 以下の判定が割れた行を、両方の値を並べて出す。"""
    ws = wb.create_sheet("要確認")
    ws["A1"] = "帳票と koutei の DB で「A3 以下」の判定が割れた行"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = (f"月シートには {taken} の値を載せている。もう一方の値もこの表に残してあるので、"
                "気になる行は帳票リンクで元を確かめる。"
                "帳票側は用紙の大きさを仕上りサイズ欄に拾ってしまうことがあり、DB 側は部品を多く並べる"
                "（付属のポスターなど別部品や、MIS 一括取込の既定値「B1」が混ざる）。")
    ws["A2"].font = FONT_NOTE
    hrow = 4
    for i, (h, w) in enumerate(CONFLICT_COLS, 1):
        ws.cell(hrow, i, h)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, hrow, 7, len(CONFLICT_COLS))
    r = hrow
    for n, (x, d, b, url) in enumerate(rows, 1):
        r += 1
        vals = [n, x.get("月", ""), x["営業部"], x["管理番号"], x["ｸﾗｲｱﾝﾄ名"], x["品名"], x["通し数"],
                d["仕上りサイズ"], d["A3以下"], b["仕上りサイズ"], b["A3以下"], taken, "開く" if url else ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border, c.alignment = FONT, BORDER, ALIGN_WRAP
            if n % 2 == 0:
                c.fill = FILL_BAND
        for i in (9, 11):
            c = ws.cell(r, i)
            if c.value in FILL_A3:
                c.fill, c.alignment = FILL_A3[c.value], ALIGN_CENTER
        ws.cell(r, 7).number_format = "#,##0"
        ws.cell(r, 7).alignment = ALIGN_RIGHT
        ws.cell(r, 12).alignment = ALIGN_CENTER
        ws.cell(r, 12).font = FONT_BOLD
        if url:
            link = ws.cell(r, len(CONFLICT_COLS))
            link.hyperlink, link.font = url, FONT_LINK
    ws.freeze_panes = ws.cell(hrow + 1, 5)
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(CONFLICT_COLS))}{max(r, hrow + 1)}"
    ws.print_title_rows = f"{hrow}:{hrow}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return r - hrow


def write_summary(wb, months, ranges, stats, base_name):
    ws = wb.create_sheet("集計", 0)
    ws["A1"] = "平版印刷 → オンデマンド移行検討（通し数 3,000 以下）　月別の集計"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = f"作成 {dt.date.today().isoformat()}　母集団: {base_name}　件数は各月シートを数えた式（月シートを直せば追従する）"
    ws["A2"].font = FONT_NOTE
    heads = ["月", "件数", "値あり", "空欄", "うち帳票", "うち DB", "うち再版推定",
             "仕上りサイズあり", "A3以下 ○", "A3以下 ×", "A3以下 不明",
             "内作", "外注", "内作・外注", "加工内容あり", "用紙銘柄あり"]
    hrow = 4
    for i, h in enumerate(heads, 1):
        ws.cell(hrow, i, h)
    style_header(ws, hrow, len(heads), len(heads))
    L = get_column_letter
    c_size, c_a3, c_cont, c_kind, c_paper = L(N_BASE + 1), L(N_BASE + 2), L(N_BASE + 3), L(N_BASE + 4), L(N_BASE + 6)
    c_src = L(N_BASE + 12)   # 「元の帳票」列。帳票名／koutei(DB)／再版推定 nnnnnnn のどれかが入る
    r = hrow
    for mth in months:
        r += 1
        first, last = ranges[mth]
        q = f"'{mth}'"
        def rng(col):
            return f"{q}!{col}{first}:{col}{last}"
        n_all, n_have = f'COUNTA({rng("C")})', f'COUNTA({rng(c_src)})'
        n_db, n_guess = f'COUNTIF({rng(c_src)},"{SRC_DB}*")', f'COUNTIF({rng(c_src)},"{SRC_GUESS}*")'
        vals = [mth, f'={n_all}',
                f'={n_have}', f'={n_all}-{n_have}',
                f'={n_have}-{n_db}-{n_guess}', f'={n_db}', f'={n_guess}',
                f'=COUNTA({rng(c_size)})',
                f'=COUNTIF({rng(c_a3)},"○")', f'=COUNTIF({rng(c_a3)},"×")', f'=COUNTIF({rng(c_a3)},"不明")',
                f'=COUNTIF({rng(c_kind)},"内作")', f'=COUNTIF({rng(c_kind)},"外注")', f'=COUNTIF({rng(c_kind)},"内作・外注")',
                f'=COUNTA({rng(c_cont)})', f'=COUNTA({rng(c_paper)})']
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = FONT, BORDER
    last = r
    r += 1
    ws.cell(r, 1, "合計").font = FONT_BOLD
    ws.cell(r, 1).border = BORDER
    for i in range(2, len(heads) + 1):
        col = L(i)
        # 月が 1 つも無いときに =SUM(B5:B4) という自分自身を含む式を作らない
        c = ws.cell(r, i, f"=SUM({col}{hrow + 1}:{col}{last})" if last > hrow else 0)
        c.font, c.border = FONT_BOLD, BORDER
    r += 2
    ws.cell(r, 1, "凡例").font = FONT_BOLD
    legend = [("○", FILL_A3["○"], "仕上りサイズが A3 以下（オンデマンド機に載る大きさ）"),
              ("×", FILL_A3["×"], "1 部品でも A3 より大きい"),
              ("不明", FILL_A3["不明"], "仕上りサイズが「規格外」だけで実寸が無い"),
              ("帳票なし", FILL_NONE, "ドライブに製造指示書などの PDF が無く、帳票の列が埋められなかった"),
              ("", FILL_DB, "koutei の DB から埋めた行（「元の帳票」列が koutei(DB)）"),
              ("", FILL_GUESS, "再版の帳票から推定した行（「元の帳票」列が 再版推定 nnnnnnn）。確かな値ではない")]
    for label, fill, text in legend:
        r += 1
        c = ws.cell(r, 1, label)
        c.fill, c.font, c.border, c.alignment = fill, FONT, BORDER, ALIGN_CENTER
        ws.cell(r, 2, text).font = FONT
    r += 2
    notes = ["この資料の作り方と注意",
             f"・母集団は「{base_name}」の明細のうち通し数 3,000 以下の {stats['rows']} 行（管理番号 {stats['keys']} 件）。月は印刷日の月",
             f"・帳票の列は Google ドライブ「作業依頼書」「委託業務依頼書」の PDF（製造指示書 {stats['seizo']}、印刷指示書 {stats['insatsu']}、外注委託依頼書 {stats['itaku']} 通）の文字を読んで管理番号で付けた",
             "・加工内容は帳票の加工欄から機械的に拾った言葉。迷うときは「加工欄の原文」列と「帳票リンク」で元を見る",
             "・備考案は 外注（委託先）：加工内容 ／ 内作（加工所）：加工内容 の形。元の Excel の備考欄に貼るときの案",
             "・製造指示書の得意先名・製品名の切れ目、用紙の銘柄の切れ目は機械的に推定している",
             "・同じ管理番号が複数行ある月は、稼動日報の実績どおり複数行のまま。帳票の列は同じ値が入る",
             "・薄紫の行は「再版推定」。その管理番号を前回受注番号に持つ再版の帳票から写した値で、帳票そのものではない",
             "　元も再版も帳票がある 473 組で確かめたところ、両方に仕上りサイズがある 469 組で A3 以下の判定が 96.4% 一致。",
             "　仕上りサイズそのものは 82%、加工内容は 67% しか一致しない（連番の別部品を前回受注番号に入れている帳票が混ざるため）",
             "　推定の行のリンクは、推定の元にした別の受注番号の帳票を開く。決める前に元の帳票か基幹システムで確かめる"]
    if stats.get("dropped"):
        notes += [f"・品名に {stats['drop_words']} を含む {stats['dropped']} 行（通し数 {stats['drop_pass']:,}）は外してある。"
                  "本刷りではないため"]
    if stats.get("blank_process"):
        notes += [f"・加工内容が空だった行には「{stats['blank_process']}」を入れてある（加工が無い＝化粧断裁だけ、という扱い）。"
                  "帳票にも DB にも行が無い空欄の行には入れていない"]
    if stats.get("filled"):
        notes += [f"・帳票のある {stats['filled']} 行は、帳票で空だった欄だけ koutei の DB の値で補ってある"
                  "（「元の帳票」列が 製造指示書・koutei(DB) のように並ぶ。帳票にある値は上書きしていない）"]
    if stats.get("conflicts"):
        notes += [f"・「要確認」シートに {stats['conflicts']} 行。帳票と koutei の DB で A3 以下の判定が割れた行で、"
                  f"月シートには {stats.get('taken', '帳票')} の値を載せている",
                  "　帳票側は用紙の大きさを仕上りサイズ欄に拾うことがあり、DB 側は別部品や MIS 一括取込の既定値「B1」を"
                  "並べることがある。どちらも取りこぼすので元を見て決める"]
    if stats.get("a3_only"):
        notes += ["・この資料は A3 以下（○）の行だけ。A3 より大きい（×）行は外してあるので、合計は母集団の件数より少ない",
                  "・仕上りサイズが分からない行は「未判定」シートにまとめてある。× かどうかはまだ決まっていない",
                  "・薄紫の「再版推定」の行もそのまま ○ 側に入っている。A3 判定の一致率は 96.4% なので、決める前に元を確かめる"]
    for i, t in enumerate(notes):
        ws.cell(r + i, 1, t).font = FONT_BOLD if i == 0 else FONT_NOTE
    for i, w in enumerate([10, 8, 8, 8, 9, 8, 11, 13, 10, 10, 11, 8, 8, 11, 12, 12], 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(hrow + 1, 2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base", required=True, help="全社印刷実績の Excel")
    ap.add_argument("--textdir", required=True, help="帳票 PDF の文字（pdftext）フォルダ")
    ap.add_argument("--out", default="平版印刷_オンデマンド移行検討_3000通し以下_資料.xlsx")
    ap.add_argument("--max-pass", type=int, default=3000)
    ap.add_argument("--csv", help="同じ内容の CSV も出す")
    ap.add_argument("--koutei-csv", help="export_finish_size.py が出した CSV。帳票が無い行をこれで埋める")
    ap.add_argument("--guess-reprint", action="store_true",
                    help="帳票も DB も無い行を、その管理番号を前回受注番号に持つ再版の帳票から推定して埋める（薄紫）")
    ap.add_argument("--conflict-size", choices=["帳票", "DB"], default="帳票",
                    help="帳票と DB で A3 の判定が割れたとき、どちらの仕上りサイズを載せるか（既定: 帳票）。"
                         "どちらを選んでも「要確認」シートに両方の値を並べる")
    ap.add_argument("--exclude-name", action="append", default=[], metavar="語",
                    help="品名にこの語を含む行を外す（何度でも指定できる）。例: --exclude-name 本機校正")
    ap.add_argument("--blank-process", metavar="語",
                    help="加工内容が空の行にこの語を入れる（加工が無い＝化粧断裁だけ、という意味のとき）。"
                         "出どころが 1 つも無い行には入れない")
    ap.add_argument("--a3-only", action="store_true",
                    help="月シートは A3 以下（○）の行だけにし、× は外す。帳票なし・不明は「未判定」シートにまとめる")
    args = ap.parse_args()

    base = read_base(args.base)
    small = [x for x in base if x["通し数"] is not None and x["通し数"] <= args.max_pass]
    print(f"母集団: 明細 {len(base)} 行 → 通し数 {args.max_pass:,} 以下 {len(small)} 行（管理番号 {len({x['key'] for x in small})} 件）")
    dropped = []
    if args.exclude_name:
        keep = []
        for x in small:
            word = next((w for w in args.exclude_name if w in (x["品名"] or "")), None)
            (dropped if word else keep).append(x)
        small = keep
        n_pass = sum(x["通し数"] or 0 for x in dropped)
        print(f"品名で外した行: {len(dropped)}（{'、'.join(args.exclude_name)}／通し数 {n_pass:,}）→ 残り {len(small)} 行")

    all_docs = load_docs(args.textdir)
    merged = merge(all_docs)
    docs = {r["key"]: dict(r, _src=SRC_DOC) for r in merged}
    n_doc = len(docs)
    n_db = n_guess = 0
    db_rows = {}
    conflicts_all = {}
    n_fill = 0
    if args.koutei_csv:          # 帳票が無い管理番号を DB で埋め、帳票にある管理番号は空欄だけ補う
        db_rows = read_koutei_csv(args.koutei_csv)
        for k, v in db_rows.items():
            d = docs.get(k)
            if d is None:
                docs[k] = v
                n_db += 1
                continue
            if d.get("_src") != SRC_DOC:
                continue
            filled = [f for f in MERGE_FIELDS if not str(d.get(f) or "").strip() and str(v.get(f) or "").strip()]
            if not filled:
                continue
            for f in filled:
                d[f] = v[f]
            if "仕上りサイズ" in filled:      # 判定もその値から出し直す
                d["A3以下"] = v["A3以下"]
            d["帳票"] = list(d["帳票"]) + [SRC_DB]   # 「元の帳票」に koutei(DB) を足して出どころを残す
            n_fill += 1
        print(f"koutei の DB: {args.koutei_csv} から {n_db} 件を補い、帳票のある {n_fill} 件の空欄を埋めた")

        # 帳票と DB で A3 の判定が割れた行。どちらを載せるかは --conflict-size で決める
        for k, v in db_rows.items():
            d = docs.get(k)
            if d is None or d.get("_src") != SRC_DOC:
                continue
            if "不明" in (d["A3以下"], v["A3以下"]) or d["A3以下"] == v["A3以下"]:
                continue
            was = {"仕上りサイズ": d["仕上りサイズ"], "A3以下": d["A3以下"]}
            if args.conflict_size == "DB":
                d["仕上りサイズ"], d["A3以下"] = v["仕上りサイズ"], v["A3以下"]
                if SRC_DB not in d["帳票"]:
                    d["帳票"] = list(d["帳票"]) + [SRC_DB]
            conflicts_all[k] = (was, {"仕上りサイズ": v["仕上りサイズ"], "A3以下": v["A3以下"]})
        if conflicts_all:
            print(f"帳票と DB で A3 の判定が割れた受注番号: {len(conflicts_all)} 件（採用: {args.conflict_size}）")
    if args.blank_process:       # 加工が無い＝化粧断裁だけ、という扱い。出どころのある行にだけ入れる
        n_bp = 0
        for d in docs.values():
            if not str(d.get("加工内容") or "").strip():
                d["加工内容"] = args.blank_process
                n_bp += 1
        print(f"加工内容が空の {n_bp} 件に「{args.blank_process}」を入れた（出どころが無い行には入れない）")

    if args.guess_reprint:       # それでも無い行を再版の帳票から推定（薄紫・確かな値ではない）
        for k, v in guess_from_reprint(merged).items():
            if k not in docs:
                docs[k] = v
                n_guess += 1
        print(f"再版元からの推定: {n_guess} 件（推定なので「元の帳票」列で見分けられるようにしてある）")
    stats = {"rows": len(small), "keys": len({x["key"] for x in small}), "a3_only": args.a3_only,
             "dropped": len(dropped), "drop_words": "・".join(args.exclude_name),
             "drop_pass": sum(x["通し数"] or 0 for x in dropped),
             "blank_process": args.blank_process, "filled": 0,
             "seizo": sum(d["doc"] == "製造指示書" for d in all_docs), "insatsu": sum(d["doc"] == "印刷指示書" for d in all_docs),
             "itaku": sum(d["doc"] == "外注委託依頼書" for d in all_docs)}
    stats["filled"] = n_fill
    hit = sum(1 for x in small if x["key"] in docs)
    src_rows = collections.Counter(docs[x["key"]].get("_src") for x in small if x["key"] in docs)
    print(f"帳票: {len(all_docs)} 通・受注 {n_doc} 件 → 母集団の埋まった行 {hit} / {len(small)}"
          f"（{'、'.join(f'{k} {v}' for k, v in sorted(src_rows.items()))}）")

    by_month = collections.defaultdict(list)
    pending = []
    for x in small:
        m = x["印刷日"][:7].replace("/", "-") or "月不明"
        if args.a3_only:
            d = docs.get(x["key"])
            verdict = d["A3以下"] if d else "帳票なし"
            if verdict == "×":
                continue
            if verdict != "○":
                x = dict(x, 月=m)
                pending.append(x)
                continue
        by_month[m].append(x)
    months = sorted(by_month)
    for m in months:
        by_month[m].sort(key=lambda x: (x["印刷日"], x["営業部"], x["管理番号"]))

    wb = Workbook()
    wb.remove(wb.active)
    ranges = {}
    note = ("左の紺色の列は稼動日報（全社印刷実績）、右の青色の列は製造指示書などの帳票 PDF から。"
            "A3以下: 緑=○／橙=×／黄=不明／灰=帳票なし。加工内容は機械的に拾った言葉なので、迷うときは右端の原文とリンク先で確かめる")
    for m in months:
        ranges[m] = write_month(wb, m, by_month[m], docs, note)
        print(f"  {m}: {len(by_month[m])} 行、埋まった {sum(1 for x in by_month[m] if x['key'] in docs)}")
    sheets = [(m, by_month[m]) for m in months]
    if args.a3_only and pending:
        pending.sort(key=lambda x: (x["月"], x["印刷日"], x["営業部"], x["管理番号"]))
        ranges["未判定"] = write_month(wb, "未判定", pending, docs,
                                     "仕上りサイズが分からない行（帳票なし、または規格外で実寸なし）。A3 以下かどうかは元の帳票か基幹システムで確かめる")
        print(f"  未判定: {len(pending)} 行")
        sheets.append(("未判定", pending))
    # 判定が割れた受注番号のうち、この資料に出てくるものを「要確認」シートに並べる
    conflicts, seen = [], set()
    for name, xs in sheets:
        for x in xs:
            k = x["key"]
            if k not in conflicts_all or k in seen:
                continue
            seen.add(k)
            was, now = conflicts_all[k]
            conflicts.append((dict(x, 月=x.get("月") or x["印刷日"][:7].replace("/", "-")),
                              was, now, (docs.get(k) or {}).get("url", "")))
    if conflicts:
        conflicts.sort(key=lambda t: (t[0]["月"], t[0]["管理番号"]))
        n_conf = write_conflicts(wb, conflicts, args.conflict_size)
        stats["conflicts"] = n_conf
        stats["taken"] = args.conflict_size
        print(f"  要確認: {n_conf} 行（帳票と DB で A3 の判定が割れた。採用: {args.conflict_size}）")
    write_summary(wb, [name for name, _ in sheets], ranges, stats, os.path.basename(args.base))
    wb.save(args.out)
    print(f"出力: {args.out}")

    if args.csv:
        import csv
        with open(args.csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, lineterminator="\r\n")
            w.writerow(["シート", "月"] + [h.replace("\n", "") for h, _ in BASE_COLS]
                       + [h.replace("\n", "") for h, _ in DOC_COLS if h != "帳票\nリンク"] + ["リンク"])
            for sheet, xs in sheets:        # 未判定シートの行も落とさない
                for x in xs:
                    d = docs.get(x["key"])
                    row = [sheet, x.get("月") or x["印刷日"][:7].replace("/", "-") or "月不明",
                           x["営業部"], x["管理番号"], x["ｸﾗｲｱﾝﾄ名"], x["品名"], x["営業担当ｺｰﾄﾞ"], x["印刷日"], x["色数"], x["通し数"]]
                    if d:
                        row += [d["仕上りサイズ"], d["A3以下"], d["加工内容"], d["内外作"], d["加工所"], d["用紙銘柄"], d["用紙規格"], d["連量"],
                                d["総頁数"], d["受注数量"], remark(d), "・".join(d["帳票"]), d["加工原文"], d["url"]]
                    else:
                        row += ["", "帳票なし"] + [""] * (len(DOC_COLS) - 2)
                    w.writerow(row)
        print(f"      {args.csv}")


if __name__ == "__main__":
    main()
