#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Google ドライブの帳票 PDF（製造指示書・印刷指示書・外注委託依頼書）から取り出した文字を読んで、
受注ごとに 仕上りサイズ・加工・用紙・通し数 をまとめ、月別シートの Excel 資料にする。

    python build_from_pdf.py --textdir <pdftext フォルダ> --out 平版印刷_オンデマンド移行検討_資料.xlsx
        [--months 2026-04 2026-09] [--max-pass 3000]

* 入力は pdftext/<fileId>.json（{"id","title","createdTime","text"}）。ドライブから文字を取り出したもの
* 1 受注に複数の帳票があるときは、製造指示書を主にし、印刷指示書・外注委託依頼書で補う
* 月は 出荷日 → 納品日 → 受注日 → 帳票の作成日 の順で決める（稼動日報の印刷月とは数日ずれることがある）
* DB は使わない。ドライブも読むだけ
"""

import argparse
import collections
import datetime as dt
import glob
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from export_finish_size import judge_a3_all, norm_order  # noqa: E402
from shiji_parse import parse_any  # noqa: E402

FONT_NAME = "Meiryo"
FONT = Font(name=FONT_NAME, size=9)
FONT_BOLD = Font(name=FONT_NAME, size=9, bold=True)
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
FONT_NOTE = Font(name=FONT_NAME, size=8, color="595959")
FONT_HEADER = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
FONT_LINK = Font(name=FONT_NAME, size=9, color="0563C1", underline="single")
FILL_HEADER = PatternFill("solid", fgColor="1F4E78")
FILL_BAND = PatternFill("solid", fgColor="F2F2F2")
FILL_A3 = {"○": PatternFill("solid", fgColor="C6EFCE"), "×": PatternFill("solid", fgColor="F8CBAD"),
           "不明": PatternFill("solid", fgColor="FFF2CC")}
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_NG = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALIGN_WRAP = Alignment(vertical="top", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

COLUMNS = [
    # (見出し, 幅)
    ("No", 5), ("受注番号", 10), ("得意先", 22), ("製品名", 28), ("所属", 11),
    ("受注数量", 9), ("通し数\n(計画)", 9), ("通し\n3000以下", 8),
    ("仕上りサイズ", 14), ("A3以下", 7), ("総頁数", 6), ("色数", 6),
    ("内外作", 7), ("加工所／委託先", 14), ("加工内容", 30),
    ("用紙銘柄", 18), ("用紙規格", 10), ("連量", 8),
    ("出荷日", 11), ("納品日", 11), ("受注日", 11),
    ("元の帳票", 14), ("製造指示書\nリンク", 12), ("加工欄の原文", 40),
]


def month_of(rec):
    for k in ("出荷日", "納品日", "受注日"):
        v = rec.get(k) or ""
        m = re.match(r"(\d{4})/(\d{1,2})", v)
        if m:
            return f"{m.group(1)}-{int(m.group(2)):02d}"
    c = rec.get("createdTime") or ""
    return c[:7] if re.match(r"\d{4}-\d{2}", c) else ""


def uniq(seq):
    out = []
    for x in seq:
        if x and x not in out:
            out.append(x)
    return out


def site_key(name):
    """委託先名の表記ゆれ（㈱松岡製本／松岡製本／八王子紙工株式会社）を同じ会社とみなすための鍵。"""
    import unicodedata
    k = unicodedata.normalize("NFKC", name)
    k = re.sub(r"(株式会社|有限会社|\(株\)|\(有\)|㈱|㈲|様)", "", k)
    return re.sub(r"\s+", "", k)


def uniq_sites(seq):
    """同じ会社は 1 つにまとめ、いちばん長い（正式に近い）表記を残す。"""
    best = {}
    order = []
    for x in seq:
        if not x:
            continue
        k = site_key(x)
        if k not in best:
            best[k] = x
            order.append(k)
        elif len(x) > len(best[k]):
            best[k] = x
    return [best[k] for k in order]


def load_docs(textdir):
    docs = []
    for p in sorted(glob.glob(os.path.join(textdir, "*.json"))):
        if os.path.basename(p).startswith("index_"):
            continue
        try:
            j = json.load(open(p, encoding="utf-8"))
        except (ValueError, OSError):
            continue
        if not isinstance(j, dict) or not j.get("text"):
            continue
        d = parse_any(j["text"], j.get("title", ""))
        if not d:
            continue
        d["fileId"] = j.get("id", "")
        d["title"] = j.get("title", "")
        d["createdTime"] = j.get("createdTime", "")
        d["url"] = f"https://drive.google.com/file/d/{j.get('id', '')}/view"
        key, _ = norm_order(d.get("受注番号") or "")
        d["key"] = key
        docs.append(d)
    return docs


def merge(docs):
    """受注番号ごとに帳票をまとめる。製造指示書 → 印刷指示書 → 委託依頼書 の順に採用する。"""
    by = collections.OrderedDict()
    for d in docs:
        if not d["key"]:
            continue
        by.setdefault(d["key"], []).append(d)
    rows = []
    for key, ds in by.items():
        seizo = sorted([d for d in ds if d["doc"] == "製造指示書"], key=lambda d: d["createdTime"])
        insatsu = sorted([d for d in ds if d["doc"] == "印刷指示書"], key=lambda d: d["createdTime"])
        itaku = sorted([d for d in ds if d["doc"] == "外注委託依頼書"], key=lambda d: d["createdTime"])
        base = (seizo or insatsu or itaku or ds)[-1]
        r = {
            "key": key,
            "受注番号": base.get("受注番号") or key,
            "得意先": base.get("得意先") or next((d.get("得意先") for d in ds if d.get("得意先")), ""),
            "製品名": base.get("製品名") or next((d.get("製品名") for d in ds if d.get("製品名")), ""),
            "所属": base.get("所属", ""),
            "受注数量": base.get("受注数量"),
            "通し数": base.get("通し数合計") if base.get("通し数合計") is not None else next((d.get("通し数合計") for d in ds if d.get("通し数合計") is not None), None),
            "仕上りサイズ": base.get("仕上りサイズ") or next((d.get("仕上りサイズ") for d in ds if d.get("仕上りサイズ")), ""),
            "総頁数": base.get("総頁数"),
            "色数": base.get("色数", ""),
            "出荷日": base.get("出荷日") or next((d.get("出荷日") for d in ds if d.get("出荷日")), ""),
            "納品日": base.get("納品日") or next((d.get("納品日") for d in ds if d.get("納品日")), ""),
            "受注日": base.get("受注日", ""),
            "createdTime": base.get("createdTime", ""),
            "帳票": uniq([d["doc"] for d in ds]),
            "url": (seizo or insatsu or ds)[-1]["url"],
            "docs": ds,
        }
        # 用紙: 製造指示書を優先。無ければ他から
        papers = base.get("用紙") or next((d.get("用紙") for d in ds if d.get("用紙")), [])
        r["用紙銘柄"] = " / ".join(uniq(p["銘柄"] for p in papers))
        r["用紙規格"] = " / ".join(uniq(p["規格"] for p in papers))
        r["連量"] = " / ".join(uniq(str(p["連量"]).rstrip("0").rstrip(".") if re.fullmatch(r"[\d.]+", str(p["連量"])) else p["連量"] for p in papers))
        # 加工: 全帳票から集める
        inout, sites, works, raw = [], [], [], []
        for d in ds:
            k = d.get("加工") or {}
            if k.get("内外"):
                inout.append(k["内外"])
            sites += k.get("加工所", [])
            works += k.get("作業", [])
            if k.get("原文"):
                raw.append(f"[{d['doc']}] {k['原文']}")
        inout = uniq(inout)
        r["内外作"] = "内作・外注" if len(inout) > 1 else (inout[0] if inout else "")
        r["加工所"] = " / ".join(uniq_sites(sites))
        r["加工内容"] = " / ".join(uniq(works))
        r["加工原文"] = "\n".join(uniq(raw))
        r["A3以下"] = judge_a3_all([r["仕上りサイズ"]]) if r["仕上りサイズ"] else "不明"
        rows.append(r)
    return rows


def style_header(ws, hrow, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(hrow, c)
        cell.font, cell.fill, cell.alignment, cell.border = FONT_HEADER, FILL_HEADER, ALIGN_CENTER, BORDER
    ws.row_dimensions[hrow].height = 30


def write_month_sheet(wb, title, rows, max_pass, note):
    ws = wb.create_sheet(title)
    ws["A1"] = f"平版印刷 → オンデマンド移行検討　{title}　（製造指示書から）"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = note
    ws["A2"].font = FONT_NOTE
    hrow = 4
    for i, (h, w) in enumerate(COLUMNS, 1):
        ws.cell(hrow, i, h)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, hrow, len(COLUMNS))
    r = hrow
    for n, x in enumerate(rows, 1):
        r += 1
        pass_ok = "" if x["通し数"] is None else ("○" if x["通し数"] <= max_pass else "×")
        vals = [n, x["受注番号"], x["得意先"], x["製品名"], x["所属"], x["受注数量"], x["通し数"], pass_ok,
                x["仕上りサイズ"], x["A3以下"], x["総頁数"], x["色数"], x["内外作"], x["加工所"], x["加工内容"],
                x["用紙銘柄"], x["用紙規格"], x["連量"], x["出荷日"], x["納品日"], x["受注日"],
                "・".join(x["帳票"]), "開く", x["加工原文"]]
        band = n % 2 == 0
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border, c.alignment = FONT, BORDER, ALIGN_WRAP
            if band:
                c.fill = FILL_BAND
        a3 = ws.cell(r, 10)
        if a3.value in FILL_A3:
            a3.fill, a3.alignment = FILL_A3[a3.value], ALIGN_CENTER
        pc = ws.cell(r, 8)
        if pc.value == "○":
            pc.fill, pc.alignment = FILL_OK, ALIGN_CENTER
        elif pc.value == "×":
            pc.fill, pc.alignment = FILL_NG, ALIGN_CENTER
        link = ws.cell(r, 23)
        link.hyperlink = x["url"]
        link.font = FONT_LINK
        for i in (6, 7, 11):
            ws.cell(r, i).number_format = "#,##0"
            ws.cell(r, i).alignment = Alignment(horizontal="right", vertical="top")
    ws.freeze_panes = ws.cell(hrow + 1, 3)
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(COLUMNS))}{max(r, hrow + 1)}"
    ws.print_title_rows = f"{hrow}:{hrow}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "&P / &N"
    return hrow + 1, r


def write_summary(wb, months, ranges, max_pass, stats):
    ws = wb.create_sheet("集計", 0)
    ws["A1"] = "平版印刷 → オンデマンド移行検討　月別の集計（製造指示書から）"
    ws["A1"].font = FONT_TITLE
    ws["A2"] = f"作成 {dt.date.today().isoformat()}　件数は各月シートを COUNTIF で数えている（月シートを直せば追従する）"
    ws["A2"].font = FONT_NOTE
    heads = ["月", "受注件数", f"通し{max_pass:,}以下", "仕上りサイズあり", "A3以下 ○", "A3以下 ×", "A3以下 不明",
             "内作", "外注", "内作・外注", "加工内容あり", "用紙銘柄あり"]
    hrow = 4
    for i, h in enumerate(heads, 1):
        ws.cell(hrow, i, h)
    style_header(ws, hrow, len(heads))
    r = hrow
    for mth in months:
        r += 1
        first, last = ranges[mth]
        q = f"'{mth}'"
        def rng(col):
            return f"{q}!{col}{first}:{col}{last}"
        vals = [mth,
                f'=COUNTA({rng("B")})',
                f'=COUNTIF({rng("H")},"○")',
                f'=COUNTA({rng("I")})',
                f'=COUNTIF({rng("J")},"○")', f'=COUNTIF({rng("J")},"×")', f'=COUNTIF({rng("J")},"不明")',
                f'=COUNTIF({rng("M")},"内作")', f'=COUNTIF({rng("M")},"外注")', f'=COUNTIF({rng("M")},"内作・外注")',
                f'=COUNTA({rng("O")})', f'=COUNTA({rng("P")})']
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = FONT, BORDER
    r += 1
    ws.cell(r, 1, "合計").font = FONT_BOLD
    ws.cell(r, 1).border = BORDER
    for i in range(2, len(heads) + 1):
        col = get_column_letter(i)
        c = ws.cell(r, i, f"=SUM({col}{hrow + 1}:{col}{r - 1})")
        c.font, c.border = FONT_BOLD, BORDER
    r += 2
    ws.cell(r, 1, "凡例").font = FONT_BOLD
    legend = [("○", FILL_A3["○"], "仕上りサイズが A3 以下（オンデマンド機に載る大きさ）／通し数が基準以下"),
              ("×", FILL_A3["×"], "1 部品でも A3 より大きい／通し数が基準より多い"),
              ("不明", FILL_A3["不明"], "仕上りサイズが無い、または「規格外」だけで実寸が無い")]
    for label, fill, text in legend:
        r += 1
        c = ws.cell(r, 1, label)
        c.fill, c.font, c.border, c.alignment = fill, FONT, BORDER, ALIGN_CENTER
        ws.cell(r, 2, text).font = FONT
    r += 2
    notes = ["この資料の作り方と注意",
             "・Google ドライブ「作業依頼書」「委託業務依頼書」フォルダの帳票 PDF（製造指示書・印刷指示書・外注委託依頼書）の文字を読んで作った",
             f"・読んだ帳票 {stats['docs']} 通（製造指示書 {stats['seizo']}、印刷指示書 {stats['insatsu']}、外注委託依頼書 {stats['itaku']}、その他 {stats['other']}）、受注 {stats['orders']} 件",
             "・月は 出荷日 → 納品日 → 受注日 → 帳票の作成日 の順で決めた。稼動日報の印刷月とは数日ずれることがある",
             "・通し数は製造指示書の計画値（部品ごとの通し数の合計）。稼動日報の実績とは違う",
             "・加工内容は帳票の加工欄から機械的に拾った言葉。判断に迷うときは「加工欄の原文」列と「製造指示書リンク」で元を見る",
             "・帳票が無い受注はこの資料に出ない。稼動日報に有って無い受注番号は、帳票が PDF で保存されていない案件",
             "・PDF の文字は表の枠が無くなって潰れているので、得意先名と製品名の切れ目、用紙の銘柄の切れ目は機械的に推定している"]
    for i, t in enumerate(notes):
        ws.cell(r + i, 1, t).font = FONT_BOLD if i == 0 else FONT_NOTE
    for i, w in enumerate([12, 10, 12, 14, 10, 10, 11, 8, 8, 11, 12, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(hrow + 1, 2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--textdir", required=True)
    ap.add_argument("--out", default="平版印刷_オンデマンド移行検討_資料.xlsx")
    ap.add_argument("--months", nargs=2, default=["2026-04", "2026-09"], help="対象の月（この範囲外は「対象外」シートへ）")
    ap.add_argument("--max-pass", type=int, default=3000, help="通し数の基準")
    ap.add_argument("--csv", help="受注ごとの一覧を CSV にも出す")
    args = ap.parse_args()

    docs = load_docs(args.textdir)
    rows = merge(docs)
    stats = {"docs": len(docs), "orders": len(rows),
             "seizo": sum(d["doc"] == "製造指示書" for d in docs), "insatsu": sum(d["doc"] == "印刷指示書" for d in docs),
             "itaku": sum(d["doc"] == "外注委託依頼書" for d in docs), "other": sum(d["doc"] == "その他" for d in docs)}
    print(f"帳票 {stats['docs']} 通 → 受注 {stats['orders']} 件")

    lo, hi = args.months
    by_month = collections.defaultdict(list)
    for r in rows:
        m = month_of(r)
        by_month[m if lo <= m <= hi else "対象外"].append(r)
    for m in by_month:
        by_month[m].sort(key=lambda r: (r["出荷日"] or "9999", r["受注番号"]))

    wb = Workbook()
    wb.remove(wb.active)
    months = sorted(k for k in by_month if k != "対象外")
    ranges = {}
    note = ("通し数と仕上りサイズは製造指示書の値。A3以下: 緑=○／橙=×／黄=不明。"
            f"通し{args.max_pass:,}以下: 緑=○／橙=×。加工内容は機械的に拾った言葉なので、迷うときは右端の原文とリンク先で確かめる")
    for m in months:
        first, last = write_month_sheet(wb, m, by_month[m], args.max_pass, note)
        ranges[m] = (first, last)
        print(f"  {m}: {len(by_month[m])} 件")
    if by_month.get("対象外"):
        write_month_sheet(wb, "対象外", by_month["対象外"], args.max_pass, "対象の月の範囲に入らない受注（参考）")
        print(f"  対象外: {len(by_month['対象外'])} 件")
    write_summary(wb, months, ranges, args.max_pass, stats)
    wb.save(args.out)
    print(f"出力: {args.out}")

    if args.csv:
        import csv
        with open(args.csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, lineterminator="\r\n")
            w.writerow(["月", "受注番号", "得意先", "製品名", "所属", "受注数量", "通し数", "仕上りサイズ", "A3以下", "総頁数", "色数",
                        "内外作", "加工所／委託先", "加工内容", "用紙銘柄", "用紙規格", "連量", "出荷日", "納品日", "受注日", "元の帳票", "リンク"])
            for m in months + (["対象外"] if by_month.get("対象外") else []):
                for x in by_month[m]:
                    w.writerow([m, x["受注番号"], x["得意先"], x["製品名"], x["所属"], x["受注数量"], x["通し数"], x["仕上りサイズ"], x["A3以下"],
                                x["総頁数"], x["色数"], x["内外作"], x["加工所"], x["加工内容"], x["用紙銘柄"], x["用紙規格"], x["連量"],
                                x["出荷日"], x["納品日"], x["受注日"], "・".join(x["帳票"]), x["url"]])
        print(f"      {args.csv}")


if __name__ == "__main__":
    main()
