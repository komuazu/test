#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_final.py の「帳票以外から埋める」経路の確認。DB もドライブも使わない。

* --koutei-csv  … export_finish_size.py の CSV から埋める（確かな値）
* --guess-reprint … 再版の帳票から推定して埋める（確かな値ではない）

どちらも「元の帳票」列と地の色で出どころが分かること、帳票が優先されること、
集計シートの内訳（帳票＋DB＋推定＝値あり、値あり＋空欄＝件数）が合うことを見る。

    python tests/test_build_final.py
"""
import csv
import glob
import json
import os
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPT = os.path.join(ROOT, "build_final.py")
sys.path.insert(0, ROOT)

import build_final as BF  # noqa: E402

HEAD = ["受注番号", "仕上りサイズ", "A3以下", "加工内容", "内外作区分", "委託先名", "用紙銘柄", "用紙規格", "斤量"]
# 母集団（全社印刷実績のダミー）。列は 営業部..通し数 の 11 列
BASE_ROWS = [
    ["本社", 8726258, "㈱テスト", "ﾊﾟﾝﾌﾚｯﾄ", "", "", "", "A01", "2026-04-10", "4+4", 2400],
    ["本社", 8726259, "㈱テスト", "ﾎﾟｽﾀｰ", "", "", "", "A01", "2026-04-11", "4+0", 800],
    ["東京", 8726260, "㈱テスト", "会社案内", "", "", "", "B02", "2026-05-12", "4+4", 1200],
    ["東京", 8726261, "㈱テスト", "ﾁﾗｼ", "", "", "", "B02", "2026-05-13", "1+0", 500],
    ["池袋", 8726262, "㈱テスト", "封筒", "", "", "", "C03", "2026-05-14", "1+0", 300],
    ["池袋", 8726263, "㈱テスト", "通し数が多い", "", "", "", "C03", "2026-05-15", "4+4", 99999],
]


def make_base(path):
    wb = Workbook()
    ws = wb.active
    heads = ["営業部", "管理番号", "ｸﾗｲｱﾝﾄ名", "品名", "今年の動向", "無しの場合の代替対策", "対策通し数",
             "営業担当ｺｰﾄﾞ", "印刷日", "色数", "通し数"]
    for i, h in enumerate(heads, 1):
        ws.cell(1, i, h)
    for r, row in enumerate(BASE_ROWS, 2):
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v)
    wb.save(path)


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(HEAD)
        w.writerows(rows)


def cols(ws, hrow=4):
    return {ws.cell(hrow, c).value: c for c in range(1, ws.max_column + 1)}


def check_unit():
    """read_koutei_csv / prev_key / guess_from_reprint を直に見る。"""
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "r.csv")
        write_csv(p, [
            ["08726258", "A4 297×210", "○", "中綴じ16P", "外注", "㈱松岡製本", "ﾕﾄﾘﾛ上質", "A全判", "70"],
            ["8726259", "", "", "", "", "", "", "", ""],            # 全部空 → 入れない
            ["8726260", "B2 728×515", "", "", "", "", "", "", ""],  # A3以下 が空 → 仕上りサイズから出す
            ["ダメな値", "A4", "○", "", "", "", "", "", ""],          # 受注番号として読めない → 飛ばす
        ])
        got = BF.read_koutei_csv(p)
        assert sorted(got) == ["8726258", "8726260"], sorted(got)
        d = got["8726258"]
        assert d["仕上りサイズ"] == "A4 297×210" and d["A3以下"] == "○"
        assert d["加工所"] == "㈱松岡製本" and d["連量"] == "70" and d["内外作"] == "外注"
        assert d["帳票"] == [BF.SRC_DB] and d["_src"] == BF.SRC_DB and d["url"] == ""
        assert got["8726260"]["A3以下"] == "×", "空の A3以下 を仕上りサイズから出せていない"
        assert BF.remark(d) == "外注（㈱松岡製本）：中綴じ16P", BF.remark(d)
    print("OK: koutei CSV の読み取り（空行・読めない番号・A3判定の補い）")

    def row(key, prev="", when="2026-01-01", **kw):
        r = dict(BF.EMPTY_DOC, key=key, 帳票=["製造指示書"], url=f"u/{key}", createdTime=when,
                 docs=[{"前回受注番号": prev}], 総頁数=8, 受注数量=100)
        r.update(kw)
        return r

    assert BF.prev_key(row("2", prev="08726258")) == "8726258", "前回受注番号を正規化できていない"
    assert BF.prev_key(row("2")) == ""
    rows = [
        row("8726300", prev="8726258", when="2026-03-01", 仕上りサイズ="A4", A3以下="○"),
        row("8726301", prev="8726258", when="2026-02-01", 仕上りサイズ="B5", A3以下="○"),  # こちらが古い
        row("8726302", prev="8726302"),                                                  # 自分自身は使わない
    ]
    g = BF.guess_from_reprint(rows)
    assert sorted(g) == ["8726258"], sorted(g)
    got = g["8726258"]
    assert got["仕上りサイズ"] == "B5", "同じ元を指す再版が複数あるとき、古い方を採っていない"
    assert got["帳票"] == ["再版推定 8726301"] and got["_src"] == BF.SRC_GUESS
    assert got["総頁数"] is None and got["受注数量"] is None, "推定に元の総頁数・受注数量を持ち込んでいる"
    assert got["url"] == "u/8726301", "元にした再版の帳票へのリンクが無い"
    print("OK: 再版元の推定（古い方を採る・自分自身を除く・件数欄を持ち込まない）")


def check_end_to_end():
    with tempfile.TemporaryDirectory() as td:
        base = os.path.join(td, "base.xlsx")
        make_base(base)
        empty = os.path.join(td, "pdftext")          # 帳票が 1 通も無い状態
        os.makedirs(empty)
        res = os.path.join(td, "koutei.csv")
        write_csv(res, [
            ["8726258", "A4 297×210", "○", "中綴じ16P", "外注", "㈱松岡製本", "ﾕﾄﾘﾛ上質", "A全判", "70"],
            ["8726260", "B2 728×515", "×", "折り", "内作", "", "", "", ""],
            ["8726261", "", "不明", "", "", "", "", "", ""],       # 全部空 → 埋めない
            ["8726263", "A4", "○", "", "", "", "", "", ""],        # 通し数 3,000 超 → 母集団に居ない
        ])
        out = os.path.join(td, "out.xlsx")
        r = subprocess.run([sys.executable, SCRIPT, "--base", base, "--textdir", empty,
                            "--koutei-csv", res, "--out", out, "--csv", os.path.join(td, "out.csv")],
                           capture_output=True, text=True)
        print(r.stdout)
        assert r.returncode == 0, r.stderr
        assert "koutei(DB) 2" in r.stdout, r.stdout

        wb = load_workbook(out)
        assert wb.sheetnames == ["集計", "2026-04", "2026-05"], wb.sheetnames
        s4 = wb["2026-04"]
        c = cols(s4)
        assert s4.cell(5, c["管理番号"]).value == "8726258"
        assert s4.cell(5, c["仕上りサイズ"]).value == "A4 297×210"
        assert s4.cell(5, c["元の帳票"]).value == BF.SRC_DB
        assert s4.cell(5, c["仕上りサイズ"]).fill.fgColor.rgb.endswith("DDEBF7"), "DB から埋めた列が水色でない"
        assert s4.cell(5, c["A3以下"]).fill.fgColor.rgb.endswith("C6EFCE"), "A3以下 は判定の色のままであるべき"
        assert s4.cell(5, c["帳票\nリンク"]).value in (None, ""), "リンクが無いのに『開く』と書いている"
        # 埋まらなかった行
        assert s4.cell(6, c["A3以下"]).value == "帳票なし" and s4.cell(6, c["元の帳票"]).value in (None, "")
        assert not s4.cell(6, c["仕上りサイズ"]).fill.fgColor.rgb.endswith("DDEBF7")
        s5 = wb["2026-05"]
        assert [s5.cell(r, c["管理番号"]).value for r in (5, 6, 7)] == ["8726260", "8726261", "8726262"]
        assert s5.cell(5, c["元の帳票"]).value == BF.SRC_DB and s5.cell(6, c["A3以下"]).value == "帳票なし"
        assert s5.max_row == 7, "通し数 3,000 超の行まで入っている"
        print("OK: koutei CSV から埋める（出どころ表示・色・空行・母集団の絞り）")

        rows = list(csv.DictReader(open(os.path.join(td, "out.csv"), encoding="utf-8-sig")))
        assert len(rows) == 5 and rows[0]["元の帳票"] == BF.SRC_DB and rows[0]["リンク"] == ""
        assert rows[1]["A3以下"] == "帳票なし"
        print("OK: CSV も同じ中身")

        recalc = glob.glob("/root/.claude/skills/synced/*/xlsx/scripts/recalc.py")
        if not recalc:
            print("recalc.py が無いので式の再計算は飛ばした")
            return
        r = subprocess.run([sys.executable, recalc[0], out, "240"], capture_output=True, text=True)
        try:
            info = json.loads(r.stdout)
        except ValueError:
            info = {"error": r.stdout + r.stderr}
        print("recalc:", info)
        assert info.get("status") == "success" and info.get("total_errors") == 0, info
        sm = load_workbook(out, data_only=True)["集計"]
        head = {sm.cell(4, i).value: i for i in range(1, sm.max_column + 1)}
        for row in (5, 6, 7):
            v = [sm.cell(row, head[h]).value for h in ("件数", "値あり", "空欄", "うち帳票", "うち DB", "うち再版推定")]
            assert v[1] + v[2] == v[0], f"値あり＋空欄 が件数と合わない: {v}"
            assert v[3] + v[4] + v[5] == v[1], f"帳票＋DB＋推定 が値ありと合わない: {v}"
        tot = {h: sm.cell(7, i).value for h, i in head.items()}
        assert (tot["件数"], tot["値あり"], tot["うち帳票"], tot["うち DB"], tot["うち再版推定"]) == (5, 2, 0, 2, 0), tot
        print("OK: 集計の内訳（帳票／DB／推定）が式で合う")


SEIZO = """【製造指示書】
金額区分 確定区分 {order} ㈱テスト {name} 260410 A01
受注数量 実内見本数 実外見本数 仕上りサイズ 総頁数
前回受注番号 {qty} 0 0 {size} 8
受注日 入稿日 下版日 2026/04/01 2026/04/02 {prev} 4+4 部品 版種
部品 内・外作 加工所 2026/04/20 2026/04/22
"""


def make_pdftext(d, docs):
    """ダミーの帳票テキスト（製造指示書）を pdftext フォルダの形で置く。"""
    os.makedirs(d, exist_ok=True)
    for i, (order, prev, size, when) in enumerate(docs):
        fid = f"id{i}"
        json.dump({"id": fid, "title": f"{order}　ﾃｽﾄ品.pdf", "createdTime": when,
                   "text": SEIZO.format(order=order, prev=prev, size=size, qty=1000, name="ﾃｽﾄ品")},
                  open(os.path.join(d, f"{fid}.json"), "w", encoding="utf-8"), ensure_ascii=False)


def check_guess_end_to_end():
    """--guess-reprint を端から端まで。再版 08726300 から元 08726258 を推定する。"""
    with tempfile.TemporaryDirectory() as td:
        base = os.path.join(td, "base.xlsx")
        make_base(base)
        textdir = os.path.join(td, "pdftext")
        make_pdftext(textdir, [
            ("08726300", "08726258", "B5 257×182", "2026-03-01T00:00:00Z"),   # 古い → こちらが採られる
            ("08726301", "08726258", "A2 594×420", "2026-04-01T00:00:00Z"),
            ("08726259", "", "A4 297×210", "2026-04-02T00:00:00Z"),           # 母集団にある帳票
        ])
        out = os.path.join(td, "g.xlsx")
        r = subprocess.run([sys.executable, SCRIPT, "--base", base, "--textdir", textdir,
                            "--guess-reprint", "--out", out], capture_output=True, text=True)
        print(r.stdout)
        assert r.returncode == 0, r.stderr
        wb = load_workbook(out)
        s4 = wb["2026-04"]
        c = cols(s4)
        row = {s4.cell(r_, c["管理番号"]).value: r_ for r_ in range(5, s4.max_row + 1)}
        g = row["8726258"]
        assert s4.cell(g, c["元の帳票"]).value == "再版推定 8726300", s4.cell(g, c["元の帳票"]).value
        assert s4.cell(g, c["仕上りサイズ"]).value == "B5 257×182", "古い方の再版を採っていない"
        assert s4.cell(g, c["A3以下"]).value == "○"
        assert s4.cell(g, c["仕上りサイズ"]).fill.fgColor.rgb.endswith("E4DFEC"), "推定の列が薄紫でない"
        assert s4.cell(g, c["帳票\nリンク"]).hyperlink is not None, "推定元の帳票へのリンクが無い"
        d = row["8726259"]
        assert s4.cell(d, c["元の帳票"]).value == "製造指示書", "帳票がある行を推定で上書きしている"
        assert not s4.cell(d, c["仕上りサイズ"]).fill.fgColor.rgb.endswith("E4DFEC")
        print("OK: 再版推定を端から端まで（古い方・薄紫・リンク・帳票優先）")


def check_a3_only_csv():
    """--a3-only のとき、未判定シートの行も CSV に出ること（出ていないと黙って消える）。"""
    with tempfile.TemporaryDirectory() as td:
        base = os.path.join(td, "base.xlsx")
        make_base(base)
        textdir = os.path.join(td, "pdftext")
        os.makedirs(textdir)
        res = os.path.join(td, "k.csv")
        write_csv(res, [
            ["8726258", "A4 297×210", "○", "", "", "", "", "", ""],
            ["8726259", "B2 728×515", "×", "", "", "", "", "", ""],      # × は外れる
            ["8726260", "規格外", "不明", "", "", "", "", "", ""],         # 未判定へ
        ])
        out, ocsv = os.path.join(td, "a.xlsx"), os.path.join(td, "a.csv")
        r = subprocess.run([sys.executable, SCRIPT, "--base", base, "--textdir", textdir,
                            "--koutei-csv", res, "--a3-only", "--out", out, "--csv", ocsv],
                           capture_output=True, text=True)
        print(r.stdout)
        assert r.returncode == 0, r.stderr
        wb = load_workbook(out)
        assert "未判定" in wb.sheetnames, wb.sheetnames
        n_x = sum(ws.max_row - 4 for ws in wb if ws.title != "集計")
        rows = list(csv.DictReader(open(ocsv, encoding="utf-8-sig")))
        assert len(rows) == n_x, f"xlsx {n_x} 行に対し CSV {len(rows)} 行（未判定が落ちている）"
        assert {r_["シート"] for r_ in rows} == {"2026-04", "未判定"}, sorted({r_["シート"] for r_ in rows})
        assert all(r_["月"].startswith("2026-0") for r_ in rows), "未判定行の月が消えている"
        assert [r_["管理番号"] for r_ in rows if r_["シート"] == "2026-04"] == ["8726258"]
        print("OK: --a3-only でも CSV に未判定シートの行が出る")


def check_input_edges():
    """こわれた入力で落ちない・黙って間違えないこと。"""
    with tempfile.TemporaryDirectory() as td:
        # 複数部品の仕上りサイズ。1 部品でも A3 より大きければ ×
        p = os.path.join(td, "multi.csv")
        write_csv(p, [["8726258", "A4 297×210 / A2 594×420", "", "", "", "", "", "", ""],
                      ["8726259", "A4 297×210 / B5 257×182", "", "", "", "", "", "", ""]])
        got = BF.read_koutei_csv(p)
        assert got["8726258"]["A3以下"] == "×", "複数部品を 1 つとして判定している"
        assert got["8726259"]["A3以下"] == "○", got["8726259"]["A3以下"]

        # A3以下 列に変な値 → 仕上りサイズから出し直す
        p = os.path.join(td, "odd.csv")
        write_csv(p, [["8726258", "A4 297×210", "はい", "", "", "", "", "", ""],
                      ["8726259", "", "OK", "中綴じ", "", "", "", "", ""]])
        got = BF.read_koutei_csv(p)
        assert got["8726258"]["A3以下"] == "○" and got["8726259"]["A3以下"] == "不明"

        # CP932（Excel で保存し直した CSV）
        p = os.path.join(td, "cp932.csv")
        with open(p, "w", encoding="cp932", newline="") as f:
            w = csv.writer(f, lineterminator="\r\n")
            w.writerow(HEAD)
            w.writerow(["8726258", "A4 297×210", "○", "中綴じ", "外注", "㈱松岡製本", "", "", ""])
        got = BF.read_koutei_csv(p)
        assert got["8726258"]["加工所"] == "㈱松岡製本", "CP932 の CSV が読めていない"

        # Excel が受け付けない制御文字
        p = os.path.join(td, "ctrl.csv")
        write_csv(p, [["8726258", "A4", "○", "折\x0b り", "", "", "", "", ""]])
        got = BF.read_koutei_csv(p)
        assert got["8726258"]["加工内容"] == "折 り", repr(got["8726258"]["加工内容"])

        # 同じ受注番号が 2 度（ゼロ埋め違い）→ 後勝ち＋警告
        p = os.path.join(td, "dup.csv")
        write_csv(p, [["8726258", "A4", "○", "", "", "", "", "", ""],
                      ["08726258", "B2 728×515", "×", "", "", "", "", "", ""]])
        got = BF.read_koutei_csv(p)
        assert len(got) == 1 and got["8726258"]["仕上りサイズ"] == "B2 728×515"

        # 見出しだけ・空・受注番号列が無い
        for rows in ([], ):
            p2 = os.path.join(td, "head.csv")
            write_csv(p2, rows)
            assert BF.read_koutei_csv(p2) == {}
        p2 = os.path.join(td, "nocol.csv")
        open(p2, "w", encoding="utf-8-sig").write("品名,備考\r\nあ,い\r\n")
        assert BF.read_koutei_csv(p2) == {}
        print("OK: 複数部品・変な A3以下・CP932・制御文字・重複・列違いの CSV")


def check_empty_population():
    """母集団が 1 行も無いとき、合計行が自分自身を指す式にならないこと。"""
    with tempfile.TemporaryDirectory() as td:
        base = os.path.join(td, "base.xlsx")
        make_base(base)
        textdir = os.path.join(td, "pdftext")
        os.makedirs(textdir)
        out = os.path.join(td, "e.xlsx")
        r = subprocess.run([sys.executable, SCRIPT, "--base", base, "--textdir", textdir,
                            "--max-pass", "0", "--out", out], capture_output=True, text=True)
        assert r.returncode == 0, r.stderr
        sm = load_workbook(out)["集計"]
        for row in range(5, sm.max_row + 1):
            if sm.cell(row, 1).value == "合計":
                vals = [sm.cell(row, i).value for i in range(2, 8)]
                assert all(v == 0 for v in vals), f"合計行に式が入っている: {vals}"
                break
        else:
            raise AssertionError("合計行が無い")
        print("OK: 母集団が空でも循環参照の式を作らない")


def main():
    check_unit()
    check_end_to_end()
    check_guess_end_to_end()
    check_a3_only_csv()
    check_input_edges()
    check_empty_population()
    print("\nALL OK")


if __name__ == "__main__":
    main()
